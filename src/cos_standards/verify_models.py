from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import yaml

from cos_standards.errors import StandardsError
from cos_standards.paths import package_repo

QUOTED_SHEET = re.compile(r"'((?:[^']|'')+)'!")
FORMULA_STYLE_MARKERS = ("formula", "link", "total", "subtotal", "major")
MODEL_KEYS = {"id", "script", "workbook", "summary_sheets"}
EXCEL_ERRORS = {
    "#DIV/0!",
    "#N/A",
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#NUM!",
    "#NULL!",
    "#GETTING_DATA",
    "#SPILL!",
    "#CALC!",
    "#BUSY!",
}
WALK_SKIP = {".git", "__pycache__", ".pytest_cache"}


def sheet_refs(formula: str, names: set[str]) -> set[str]:
    found = set()

    def take(match: re.Match) -> str:
        name = match.group(1).replace("''", "'")
        if name in names:
            found.add(name)
        return " "

    rest = QUOTED_SHEET.sub(take, formula)
    for name in names:
        if re.search(rf"(?<![\w']){re.escape(name)}!", rest):
            found.add(name)
    return found


def unreached(outbound: dict[str, set[str]], roots: list[str], indexes: list[str] | None = None) -> list[str]:
    index_set = set(indexes or [])
    seen = set()
    stack = [name for name in roots if name not in index_set]
    while stack:
        cur = stack.pop()
        if cur in seen:
            continue
        seen.add(cur)
        if cur in index_set:
            continue
        stack.extend(outbound.get(cur, ()))
    seen.update(index_set)
    return sorted(name for name in outbound if name not in seen)


def is_hardcoded_summary(value, style, row: int, prefixes: tuple[str, ...] | None = None) -> bool:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return False
    if value == 0 or row < 7:
        return False
    if not isinstance(style, str):
        return False
    markers = tuple(prefixes) if prefixes is not None else FORMULA_STYLE_MARKERS
    return style_has_marker(style, markers)


def style_has_marker(style: str, markers: tuple[str, ...]) -> bool:
    for marker in markers:
        if re.search(rf"(?i)(?<![A-Za-z0-9]){re.escape(str(marker))}(?![A-Za-z0-9])", style):
            return True
    return False


def style_markers(extra: list[str] | None = None) -> tuple[str, ...]:
    seen: list[str] = []
    for item in list(FORMULA_STYLE_MARKERS) + list(extra or []):
        text = str(item).strip()
        if not text:
            raise StandardsError("formula_style_prefixes entries must be non-empty strings")
        key = text.casefold()
        if key not in seen:
            seen.append(key)
    return tuple(seen)


def norm(value):
    if isinstance(value, bool):
        return ("bool", value)
    if isinstance(value, (int, float)):
        return ("num", float(value))
    if isinstance(value, dt.datetime):
        return ("dt", value.isoformat())
    if isinstance(value, dt.date):
        return ("d", value.isoformat())
    if isinstance(value, dt.time):
        return ("t", value.isoformat())
    if isinstance(value, str):
        return ("s", value)
    if value is None:
        return None
    return ("repr", repr(value))


def _coord_key(coord: str):
    match = re.match(r"([A-Z]+)(\d+)$", coord)
    if not match:
        return (0, 0, coord)
    letters, row = match.group(1), int(match.group(2))
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return (row, n, coord)


def dump_values(path: Path) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False)
    cells = {}
    for ws in wb.worksheets:
        sheet = {}
        for row in ws.iter_rows():
            for cell in row:
                got = norm(cell.value)
                if got is not None:
                    sheet[cell.coordinate] = got
        cells[ws.title] = sheet
    sheets = list(wb.sheetnames)
    wb.close()
    return {"sheets": sheets, "cells": cells}


def diff_values(before: dict, after: dict, limit: int = 25) -> list[str]:
    diffs = []
    if before["sheets"] != after["sheets"]:
        diffs.append(f"sheets {before['sheets']} -> {after['sheets']}")
    names = sorted(set(before["cells"]) | set(after["cells"]))
    for name in names:
        old = before["cells"].get(name, {})
        new = after["cells"].get(name, {})
        for coord in sorted(set(old) | set(new), key=_coord_key):
            if old.get(coord) != new.get(coord):
                diffs.append(f"{name}!{coord}: {old.get(coord)!r} -> {new.get(coord)!r}")
                if len(diffs) >= limit:
                    return diffs
    return diffs


def reference_graph(path: Path) -> dict[str, set[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False, read_only=True)
    names = set(wb.sheetnames)
    outbound = {name: set() for name in wb.sheetnames}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    for other in sheet_refs(cell.value, names):
                        if other != ws.title:
                            outbound[ws.title].add(other)
    wb.close()
    return outbound


def hardcoded_hits(path: Path, sheets: list[str], prefixes: tuple[str, ...]) -> list[str]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False)
    hits = []
    missing = [name for name in sheets if name not in wb.sheetnames]
    if missing:
        wb.close()
        return [f"missing summary sheet: {', '.join(missing)}"]
    for name in sheets:
        ws = wb[name]
        for row in ws.iter_rows(min_row=7):
            for cell in row:
                style = cell.style if isinstance(cell.style, str) else ""
                if is_hardcoded_summary(cell.value, style, cell.row, prefixes):
                    hits.append(f"{name}!{cell.coordinate}={cell.value} style={style}")
    wb.close()
    return hits


def parse_models(data: object) -> list[dict]:
    models = data.get("models") if isinstance(data, dict) else None
    if not isinstance(models, list):
        raise StandardsError("models file needs a models list")
    parsed = []
    for item in models:
        if not isinstance(item, dict):
            raise StandardsError("model entry must be a mapping")
        unknown = sorted(set(item) - MODEL_KEYS)
        if unknown:
            raise StandardsError(f"model {item.get('id', '?')}: unknown keys {unknown}")
        missing = sorted(MODEL_KEYS - set(item))
        if missing:
            raise StandardsError(f"model {item.get('id', '?')}: missing {missing}")
        for key in ("id", "script", "workbook"):
            if not isinstance(item[key], str) or not item[key].strip():
                raise StandardsError(f"model {item.get('id', '?')}: {key} must be text")
        sheets = item["summary_sheets"]
        if (
            not isinstance(sheets, list)
            or not sheets
            or not all(isinstance(name, str) and name.strip() for name in sheets)
        ):
            raise StandardsError(f"{item['id']}: summary_sheets must name at least one sheet")
        parsed.append(item)
    return parsed


def load_models(path: Path) -> list[dict]:
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except Exception as exc:
        raise StandardsError(f"{path}: {exc}") from exc
    return parse_models(data)


def unlisted_scripts(root: Path, models: list[dict]) -> list[str]:
    listed = {item["script"] for item in models}
    found = []
    for dirpath, dirs, files in os.walk(root):
        dirs[:] = sorted(name for name in dirs if name not in WALK_SKIP and not name.endswith(".egg-info"))
        for name in sorted(files):
            if not (name.startswith("build_") and name.endswith(".py")):
                continue
            path = Path(dirpath) / name
            rel = path.relative_to(root).as_posix()
            if rel.startswith("Archive/") or "/outputs/" in f"/{rel}":
                continue
            if rel not in listed:
                found.append(rel)
    return sorted(found)


def _fail(checks: dict, name: str, detail) -> None:
    checks[name] = "fail"
    checks[f"{name}_detail"] = detail


def recalculate(workbook: Path) -> dict:
    import openpyxl

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice is not installed (soffice or libreoffice)")
    src = workbook.resolve()
    with tempfile.TemporaryDirectory(prefix="cos-verify-") as tmp:
        tmp_path = Path(tmp)
        profile = tmp_path / "profile"
        profile.mkdir()
        outdir = tmp_path / "out"
        outdir.mkdir()
        env = os.environ.copy()
        env["SAL_USE_VCLPLUGIN"] = "svp"
        cmd = [
            soffice,
            "--headless",
            "--norestore",
            "--nolockcheck",
            "--nofirststartwizard",
            "--nologo",
            f"-env:UserInstallation=file://{profile}",
            "--calc",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(outdir),
            str(src),
        ]
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=180, env=env)
        if proc.returncode != 0:
            tail = (proc.stderr or proc.stdout or "LibreOffice failed").strip().splitlines()
            raise RuntimeError(" ".join(tail[-8:]))
        converted = outdir / src.name
        if not converted.is_file():
            produced = list(outdir.glob("*.xlsx"))
            if len(produced) != 1:
                raise RuntimeError("LibreOffice did not write a workbook")
            converted = produced[0]
        formula_wb = openpyxl.load_workbook(src, data_only=False)
        cached_wb = openpyxl.load_workbook(converted, data_only=True)
        errors = []
        checks = []
        uncached = []
        formula_cells = 0
        cached_values = 0
        for ws in formula_wb.worksheets:
            cached_ws = cached_wb[ws.title] if ws.title in cached_wb.sheetnames else None
            for row in ws.iter_rows():
                for cell in row:
                    if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                        continue
                    formula_cells += 1
                    cached = None if cached_ws is None else cached_ws[cell.coordinate].value
                    if cached is None:
                        uncached.append(f"{ws.title}!{cell.coordinate}")
                        continue
                    cached_values += 1
                    if isinstance(cached, str) and cached in EXCEL_ERRORS:
                        errors.append(f"{ws.title}!{cell.coordinate}={cached}")
                    elif str(cached) == "CHECK":
                        checks.append(f"{ws.title}!{cell.coordinate}={cached!r}")
        formula_wb.close()
        cached_wb.close()
        return {
            "formula_cells": formula_cells,
            "cached_values": cached_values,
            "errors": errors,
            "checks": checks,
            "uncached": uncached,
        }


def check_model(model: dict, skip_recalc: bool, root: Path, prefixes: tuple[str, ...]) -> dict:
    checks: dict = {}
    script = root / model["script"]
    workbook = root / model["workbook"]
    result = {
        "id": model["id"],
        "script": model["script"],
        "workbook": model["workbook"],
        "checks": checks,
    }
    if not script.is_file():
        _fail(checks, "rebuild", "build script missing")
        result["ok"] = False
        return result
    if not workbook.is_file():
        _fail(checks, "rebuild", "committed workbook missing")
        result["ok"] = False
        return result

    before = dump_values(workbook)
    proc = subprocess.run(
        [sys.executable, str(script)],
        cwd=script.parent,
        capture_output=True,
        text=True,
        timeout=600,
    )
    if proc.returncode != 0:
        tail = (proc.stderr or proc.stdout or "build failed").strip().splitlines()
        _fail(checks, "rebuild", tail[-15:])
        result["ok"] = False
        return result
    checks["rebuild"] = "pass"

    after = dump_values(workbook)
    drifted = diff_values(before, after)
    if drifted:
        _fail(checks, "drift", drifted)
    else:
        checks["drift"] = "pass"

    summaries = list(model["summary_sheets"])
    graph = reference_graph(workbook)
    orphan_problems = []
    missing = [name for name in summaries if name not in graph]
    if missing:
        orphan_problems.append(f"missing summary sheet: {', '.join(missing)}")
    orphans = unreached(graph, summaries)
    if orphans:
        orphan_problems.append("sheet is not reached from a summary sheet: " + ", ".join(orphans))
    if orphan_problems:
        _fail(checks, "orphans", orphan_problems)
    else:
        checks["orphans"] = "pass"

    hits = hardcoded_hits(workbook, summaries, prefixes)
    if hits:
        _fail(checks, "hardcoded_summary", hits[:25])
    else:
        checks["hardcoded_summary"] = "pass"

    if skip_recalc:
        checks["recalc"] = "skipped"
        checks["excel_errors"] = "skipped"
        checks["check_cells"] = "skipped"
        checks["cached_values"] = "skipped"
    else:
        try:
            rec = recalculate(workbook)
        except Exception as exc:
            _fail(checks, "recalc", str(exc))
            result["ok"] = False
            return result
        checks["recalc"] = "pass"
        result["formula_cells"] = rec["formula_cells"]
        result["cached_values"] = rec["cached_values"]
        if rec["errors"]:
            _fail(checks, "excel_errors", rec["errors"][:25])
        else:
            checks["excel_errors"] = "pass"
        if rec["checks"]:
            _fail(checks, "check_cells", rec["checks"][:25])
        else:
            checks["check_cells"] = "pass"
        if rec.get("uncached"):
            _fail(checks, "cached_values", rec["uncached"][:40])
        else:
            checks["cached_values"] = "pass"

    result["ok"] = not any(value == "fail" for value in checks.values())
    return result


def main(argv: list[str] | None = None) -> int:
    from cos_standards.config import load_config

    ap = argparse.ArgumentParser(description="Rebuild workbooks and write a receipt.")
    ap.add_argument("--receipt", default="verify-receipt.json")
    ap.add_argument("--skip-recalc", action="store_true")
    ap.add_argument("--only", action="append", default=[], help="model id to run (repeatable)")
    ap.add_argument("--config", default=None)
    ap.add_argument("--root", default=None)
    args = ap.parse_args(argv)

    root = Path(args.root).resolve() if args.root else Path.cwd().resolve()
    receipt: dict = {"ok": False, "models": [], "unlisted_scripts": []}
    code = 1
    try:
        config = load_config(args.config, root)
        if config.models_path is not None and config.models_path.is_file():
            models = load_models(config.models_path)
        else:
            models = []
        missing = unlisted_scripts(root, models)
        receipt["unlisted_scripts"] = missing
        chosen = [item for item in models if not args.only or item["id"] in args.only]
        unknown = sorted(set(args.only) - {item["id"] for item in models})
        if unknown:
            receipt["unknown_ids"] = unknown
        prefixes = style_markers(config.formula_style_prefixes)
        for model in chosen:
            print(f"verify {model['id']}", flush=True)
            receipt["models"].append(check_model(model, args.skip_recalc, root, prefixes))
        failed = [item for item in receipt["models"] if not item.get("ok")]
        receipt["ok"] = not failed and not missing and not unknown
        code = 0 if receipt["ok"] else 1
    except Exception as exc:
        receipt["error"] = str(exc)
        print(f"verify failed: {exc}", file=sys.stderr)
        code = 1
    path = Path(args.receipt)
    if not path.is_absolute():
        path = root / path
    path.write_text(json.dumps(receipt, indent=2) + "\n", encoding="utf-8")
    print(f"receipt {path} ok={receipt['ok']}")
    return code


if __name__ == "__main__":
    sys.exit(main())
