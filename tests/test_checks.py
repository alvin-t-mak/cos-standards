from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from cos_standards import (
    check_corrections,
    check_lookups,
    check_outward,
    check_rules,
    lint_format,
    lint_shape,
    verify_models,
)

ROOT = Path(__file__).resolve().parents[1]


class ShapeTests(unittest.TestCase):
    def test_five_lines_pass(self):
        text = "\n".join(f"line {i}" for i in range(1, 6))
        self.assertEqual(lint_shape.problems_for("deliverables/slack/note.md", text), [])

    def test_six_lines_fail(self):
        text = "\n".join(f"line {i}" for i in range(1, 7))
        problems = lint_shape.problems_for("deliverables/slack/note.md", text)
        self.assertEqual(len(problems), 1)
        self.assertIn("5", problems[0])

    def test_comments_do_not_count(self):
        text = "<!-- note -->\n" + "\n".join(f"line {i}" for i in range(1, 6))
        self.assertEqual(lint_shape.problems_for("deliverables/slack/note.md", text), [])

    def test_decision_shape_passes(self):
        text = "\n".join([
            "Pick a CRM by Friday. If no answer, I go with 1.",
            "1. Stay on the current tool (recommended)",
            "2. Move to the other tool",
            "Why: the move costs a quarter and the scores are close.",
            "Extra line one",
            "Extra line two",
        ])
        self.assertEqual(lint_shape.problems_for("deliverables/slack/note.md", text), [])

    def test_decision_without_deadline_fails(self):
        text = "\n".join([
            "Pick a CRM.",
            "1. Stay (recommended)",
            "2. Move",
            "Why: the scores are close.",
        ])
        problems = lint_shape.problems_for("deliverables/slack/decision.md", text)
        self.assertTrue(any("deadline" in item for item in problems))

    def test_fixture_files(self):
        good = (ROOT / "tests" / "slack" / "good-decision.md").read_text(encoding="utf-8")
        bad = (ROOT / "tests" / "slack" / "bad-decision.md").read_text(encoding="utf-8")
        self.assertEqual(lint_shape.problems_for("standards/tests/slack/good-decision.md", good), [])
        problems = lint_shape.problems_for("standards/tests/slack/bad-decision.md", bad)
        self.assertTrue(any("deadline" in item for item in problems))

    def test_decision_marker_without_shape_fails(self):
        text = "<!-- decision-request -->\nShip the note.\n"
        problems = lint_shape.problems_for("deliverables/slack/note.md", text)
        self.assertGreaterEqual(len(problems), 2)

    def test_two_recommended_fails(self):
        text = "\n".join([
            "1. Stay (recommended)",
            "2. Move (recommended)",
            "Why: both look fine.",
        ])
        problems = lint_shape.problems_for("deliverables/slack/decision.md", text)
        self.assertTrue(any("exactly one" in item for item in problems))

    def test_path_decision_is_a_request(self):
        text = "Just a short note.\n"
        problems = lint_shape.problems_for("deliverables/slack/decision.md", text)
        self.assertTrue(problems)

    def test_filename_pattern_is_the_draft(self):
        self.assertFalse(lint_shape.is_draft("GTM/note.md"))
        self.assertFalse(lint_shape.is_draft("deliverables/slack/note.md"))
        self.assertTrue(lint_shape.is_draft("deliverables/slack/note.slack.md"))
        self.assertTrue(lint_shape.is_draft("notes/update.slack.md"))

    def test_semicolon_fails(self):
        problems = lint_shape.problems_for("note.slack.md", "Shipped the note; the rest can wait.\n")
        self.assertTrue(any("semicolon" in item for item in problems))

    def test_star_bullet_fails(self):
        problems = lint_shape.problems_for("note.slack.md", "* Shipped the note\n")
        self.assertTrue(any("bullets" in item for item in problems))

    def test_deadline_without_default_fails(self):
        text = "\n".join([
            "Pick a CRM by Friday.",
            "1. Stay (recommended)",
            "2. Move",
            "Why: the scores are close.",
        ])
        problems = lint_shape.problems_for("note.slack.md", text)
        self.assertTrue(any("default-deadline" in item for item in problems))

    def test_stricter_line_cap(self):
        text = "\n".join(f"line {i}" for i in range(1, 5))
        problems = lint_shape.problems_for("note.slack.md", text, max_lines=3)
        self.assertEqual(len(problems), 1)
        self.assertIn("3", problems[0])


class CorrectionsTests(unittest.TestCase):
    def test_empty_passes(self):
        self.assertEqual(check_corrections.problems({"corrections": []}), [])

    def test_once_without_check_passes(self):
        item = {"id": "dates", "note": "Include the weekday.", "count": 1}
        self.assertEqual(check_corrections.problems({"corrections": [item]}), [])

    def test_twice_without_check_fails(self):
        item = {"id": "dates", "note": "Include the weekday.", "count": 2}
        problems = check_corrections.problems({"corrections": [item]})
        self.assertEqual(len(problems), 1)
        self.assertIn("dates", problems[0])
        self.assertIn("seen 2 times", problems[0])

    def test_twice_with_check_passes(self):
        item = {"id": "dates", "note": "Include the weekday.", "count": 2, "check": "writing-lint"}
        self.assertEqual(check_corrections.problems({"corrections": [item]}), [])

    def test_missing_file_adds_nothing(self):
        self.assertEqual(check_corrections.check_file(Path("/tmp/does-not-exist-corrections.yaml")), [])


class RulesTests(unittest.TestCase):
    def test_repo_ledger_matches_markers(self):
        self.assertEqual(check_rules.problems(check_rules.marker_ids(), check_rules.ledger()), [])

    def test_missing_marker_fails(self):
        rules = [{"id": "made-up", "uncheckable": "test"}]
        problems = check_rules.problems({}, rules)
        self.assertTrue(any("not marked" in item for item in problems))

    def test_unknown_check_fails(self):
        rules = [{"id": "made-up", "check": "not-a-check"}]
        problems = check_rules.problems({"made-up": ["AGENTS.md:1"]}, rules)
        self.assertTrue(any("not implemented" in item for item in problems))

    def test_consumer_shared_marker_needs_no_local_row(self):
        shared = [{"id": "task-dates", "check": "task-dates"}]
        markers = {"task-dates": ["AGENTS.md:1"]}
        self.assertEqual(check_rules.consumer_problems(markers, [], shared), [])

    def test_consumer_new_marker_needs_a_local_row(self):
        shared = [{"id": "task-dates", "check": "task-dates"}]
        markers = {"local-only": ["AGENTS.md:4"]}
        problems = check_rules.consumer_problems(markers, [], shared)
        self.assertTrue(any("missing from the ledger" in item for item in problems))

    def test_local_id_cannot_replace_a_shared_rule(self):
        shared = [{"id": "task-dates", "check": "task-dates"}]
        local = [{"id": "task-dates", "uncheckable": "skip it"}]
        problems = check_rules.consumer_problems({}, local, shared)
        self.assertTrue(any("collides" in item for item in problems))


class JobCheckTests(unittest.TestCase):
    def _repo(self, workflow: str, rules: str, marker: str | None = "local-export") -> Path:
        root = Path(self.enterContext(tempfile.TemporaryDirectory()))
        if marker:
            (root / "AGENTS.md").write_text(f"<!-- rule: {marker} -->\n", encoding="utf-8")
        standards = root / "standards"
        standards.mkdir()
        (standards / "rules.yaml").write_text(rules, encoding="utf-8")
        workflows = root / ".github" / "workflows"
        workflows.mkdir(parents=True)
        (workflows / "standards.yml").write_text(workflow, encoding="utf-8")
        return root

    def test_job_that_exists_passes(self):
        root = self._repo(
            "jobs:\n  generated:\n    name: Rebuild\n    runs-on: ubuntu-latest\n",
            "rules:\n  - id: local-export\n    check: job:generated\n",
        )
        self.assertEqual(check_rules.check(root, root / "standards" / "rules.yaml"), [])

    def test_missing_job_fails(self):
        root = self._repo(
            "jobs:\n  other:\n    name: generated\n    runs-on: ubuntu-latest\n",
            "rules:\n  - id: local-export\n    check: job:generated\n",
        )
        problems = check_rules.check(root, root / "standards" / "rules.yaml")
        self.assertTrue(any("job generated is not in .github/workflows" in item for item in problems))

    def test_job_check_on_a_shared_rule_id_is_rejected(self):
        root = self._repo(
            "jobs:\n  generated:\n    runs-on: ubuntu-latest\n",
            "rules:\n  - id: task-dates\n    check: job:generated\n",
            marker=None,
        )
        problems = check_rules.check(root, root / "standards" / "rules.yaml")
        self.assertTrue(any("task-dates" in item and "collides" in item for item in problems))

    def test_job_check_in_the_shared_ledger_is_rejected(self):
        problems = check_rules.problems(
            {"local-export": ["AGENTS.md:1"]},
            [{"id": "local-export", "check": "job:generated"}],
            job_ids={"generated"},
        )
        self.assertTrue(any("shared ledger cannot use a job check" in item for item in problems))


class SheetTests(unittest.TestCase):
    def test_shorter_name_does_not_match_inside_longer_name(self):
        names = {"Data", "GL Data", "Actuals", "East_Actuals"}
        refs = verify_models.sheet_refs("=A1+'GL Data'!B7+East_Actuals!C3", names)
        self.assertEqual(refs, {"GL Data", "East_Actuals"})

    def test_unreached_helper_fails(self):
        outbound = {"Summary": {"Calc"}, "Calc": set(), "Orphan": {"Helper"}, "Helper": set()}
        self.assertEqual(verify_models.unreached(outbound, ["Summary"]), ["Helper", "Orphan"])

    def test_index_tab_does_not_cover_what_only_it_reads(self):
        outbound = {"Index": {"Hidden"}, "Summary": {"Calc"}, "Hidden": set(), "Calc": set()}
        self.assertEqual(verify_models.unreached(outbound, ["Summary"], ["Index"]), ["Hidden"])

    def test_hardcoded_scope(self):
        self.assertFalse(verify_models.is_hardcoded_summary(0, "Formula", 10))
        self.assertFalse(verify_models.is_hardcoded_summary(4, "Formula", 6))
        self.assertFalse(verify_models.is_hardcoded_summary(4, "Input", 10))
        self.assertTrue(verify_models.is_hardcoded_summary(4, "Formula $", 8))
        self.assertTrue(verify_models.is_hardcoded_summary(4, "Brand Formula $", 8))
        self.assertFalse(verify_models.is_hardcoded_summary(4, "Hyperlink", 8))
        self.assertTrue(verify_models.is_hardcoded_summary(4, "Link", 8))
        self.assertFalse(verify_models.style_has_marker("Subtotal", ("total",)))
        self.assertTrue(verify_models.style_has_marker("Subtotal", ("subtotal",)))
        self.assertFalse(verify_models.is_hardcoded_summary(True, "Formula", 8))

    def test_model_allow_is_rejected(self):
        data = {
            "models": [
                {
                    "id": "demo",
                    "script": "build_demo.py",
                    "workbook": "demo.xlsx",
                    "summary_sheets": ["Summary"],
                    "allow": [],
                }
            ]
        }
        with self.assertRaises(Exception) as caught:
            verify_models.parse_models(data)
        self.assertIn("allow", str(caught.exception))

    def test_empty_summary_sheets_fail(self):
        data = {
            "models": [
                {
                    "id": "demo",
                    "script": "build_demo.py",
                    "workbook": "demo.xlsx",
                    "summary_sheets": [],
                }
            ]
        }
        with self.assertRaises(Exception):
            verify_models.parse_models(data)


class FormatTests(unittest.TestCase):
    def test_valid_task_date(self):
        text = "- [ ] Send it 📅 2026-08-21\n"
        self.assertEqual(lint_format.task_date_problems("Tasks.md", text), [])

    def test_prose_emoji_is_not_a_date(self):
        text = "- [ ] It was wearing a false 📅) ⏳ 2026-08-21\n"
        self.assertEqual(lint_format.task_date_problems("Tasks.md", text), [])

    def test_malformed_task_date(self):
        text = "- [ ] Send it 📅 2026-13-01\n"
        self.assertTrue(lint_format.task_date_problems("Tasks.md", text))

    def test_atomic_close_passes(self):
        text = "\n".join([
            "---",
            "status: closed",
            "closed: 2026-08-17",
            "---",
            "",
            "## Status",
            "",
            "CLOSED: shipped, 2026-08-17.",
            "",
            "- [x] Shipped ✅ 2026-08-17",
        ])
        self.assertEqual(lint_format.atomic_close_problems("note.md", text), [])

    def test_atomic_close_partial_fails(self):
        text = "---\nstatus: closed\n---\n\n## Status\n\nStill working.\n"
        problems = lint_format.atomic_close_problems("note.md", text)
        self.assertGreaterEqual(len(problems), 2)

    def test_vendor_caps(self):
        text = "### CRM\n" + "\n".join(f"- point {i} " + ("word " * 10) for i in range(5))
        problems = lint_format.vendor_notes_problems(text, "Vendors/Vendor Notes.md")
        self.assertTrue(any("bullets" in item for item in problems))

    def test_vendor_date_stamp(self):
        text = "### CRM\n- Flagged on 2026-05-11 as the pick.\n"
        problems = lint_format.vendor_notes_problems(text, "Vendors/Vendor Notes.md")
        self.assertTrue(any("date" in item for item in problems))

    def test_notion_day_label(self):
        bad = lint_format.notion_day_problems("Product/plan.md", "Monday: light\n", ("Product/",))
        ok = lint_format.notion_day_problems(
            "Product/plan.md", "Feedback arrived over the weekend.\n", ("Product/",)
        )
        outside = lint_format.notion_day_problems("GTM/note.md", "Monday: buffer\n", ("Product/",))
        self.assertTrue(bad)
        self.assertEqual(ok, [])
        self.assertEqual(outside, [])

    def test_empty_notion_paths_keep_the_default_scope(self):
        self.assertTrue(lint_format.notion_day_problems("Product/plan.md", "Monday: light\n"))
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            note = root / "Product" / "plan.md"
            note.parent.mkdir()
            note.write_text("Monday: light\n", encoding="utf-8")
            self.assertTrue(lint_format.problems(root, notion_paths=[]))

    def test_repo_format_is_clean(self):
        self.assertEqual(lint_format.problems(), [])


class LookupTests(unittest.TestCase):
    def test_xlookup_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "build_demo.py"
            path.write_text('value = "XLOOKUP(A1,B:B,C:C)"\n', encoding="utf-8")
            self.assertTrue(check_lookups.problems(Path(tmp)))

    def test_index_match_passes(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "build_demo.py"
            path.write_text('value = "INDEX(C:C, MATCH(A1, B:B, 0))"\n', encoding="utf-8")
            self.assertEqual(check_lookups.problems(Path(tmp)), [])


class ModelRunTests(unittest.TestCase):
    def test_unreferenced_sheet_fails(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Summary"
            ws["A1"] = 1
            wb.create_sheet("Orphan")
            wb.save(root / "book.xlsx")
            wb.close()
            (root / "build_book.py").write_text("print('ok')\n", encoding="utf-8")
            model = {
                "id": "demo",
                "script": "build_book.py",
                "workbook": "book.xlsx",
                "summary_sheets": ["Summary"],
            }
            result = verify_models.check_model(model, True, root, verify_models.FORMULA_STYLE_MARKERS)
        self.assertEqual(result["checks"]["rebuild"], "pass")
        self.assertEqual(result["checks"]["drift"], "pass")
        self.assertEqual(result["checks"]["orphans"], "fail")
        self.assertTrue(any("Orphan" in item for item in result["checks"]["orphans_detail"]))

    def test_recalculate_flags_check_values_and_excel_errors(self):
        import shutil

        import openpyxl

        if not shutil.which("soffice") and not shutil.which("libreoffice"):
            self.skipTest("LibreOffice is not installed")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Summary"
            ws["A1"] = "CHECK"
            ws["C7"] = 10
            ws["C8"] = 10
            ws["D7"] = '=IF(ABS(C7-C8)<1,"-","CHECK")'
            ws["C9"] = 10
            ws["C10"] = 20
            ws["D9"] = '=IF(ABS(C9-C10)<1,"-","CHECK")'
            ws["E9"] = '=IF(ABS(C9-C10)<1,"-","FAIL")'
            ws["D2"] = "=1/0"
            wb.save(path)
            wb.close()
            rec = verify_models.recalculate(path)
        self.assertEqual(rec["uncached"], [])
        self.assertTrue(any("D2=#DIV/0!" in item for item in rec["errors"]))
        joined = " ".join(rec["checks"])
        self.assertNotIn("A1", joined)
        self.assertNotIn("D7", joined)
        self.assertNotIn("E9", joined)
        self.assertIn("D9", joined)
        self.assertIn("CHECK", joined)
        self.assertNotIn("FAIL", joined)


class OutwardTests(unittest.TestCase):
    def test_term_hit_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            deck = root / "GTM"
            deck.mkdir()
            (deck / "note.md").write_text("We use zz-local-term here.\n", encoding="utf-8")
            problems = check_outward.problems({"folders": ["GTM"], "terms": ["zz-local-term"]}, root)
        self.assertTrue(problems)

    def test_other_words_pass(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            deck = root / "GTM"
            deck.mkdir()
            (deck / "note.md").write_text("The launch date is Friday.\n", encoding="utf-8")
            problems = check_outward.problems({"folders": ["GTM"], "terms": ["zz-local-term"]}, root)
        self.assertEqual(problems, [])

    def test_missing_folder_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            problems = check_outward.problems({"folders": ["Missing"], "terms": ["acme"]}, root)
        self.assertTrue(any("missing outward folder: Missing" in item for item in problems))

    def test_missing_folder_fails_when_terms_are_empty(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            problems = check_outward.problems({"folders": ["Missing"], "terms": []}, root)
        self.assertEqual(problems, ["missing outward folder: Missing"])

    def test_every_matching_line_is_reported(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for name in ("tests", "build", "dist", ".venv", "node_modules"):
                folder = root / name
                folder.mkdir()
                (folder / "note.md").write_text("acme one\nacme two\n", encoding="utf-8")
            problems = check_outward.problems(
                {"folders": ["tests", "build", "dist", ".venv", "node_modules"], "terms": ["acme"]},
                root,
            )
        self.assertEqual(len(problems), 10)
        self.assertTrue(any(item.startswith("tests/note.md:") for item in problems))

    def test_allow_key_fails(self):
        problems = check_outward.problems({"folders": ["GTM"], "terms": ["x"], "allow": []}, Path("."))
        self.assertTrue(any("allow" in item for item in problems))


if __name__ == "__main__":
    unittest.main(verbosity=1)
